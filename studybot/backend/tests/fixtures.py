"""Builders for small synthetic PDF / PPTX / DOCX files."""
from pathlib import Path

import pymupdf
from docx import Document
from pptx import Presentation


def make_pdf(path: Path, pages: list[list[tuple[str, str, tuple]]]) -> Path:
    """pages: list of pages; each page is a list of (style, text, rect) with style in
    title|h1|h2|body|footer."""
    fonts = {"title": ("hebo", 20), "h1": ("hebo", 14), "h2": ("hebo", 13),
             "body": ("helv", 10), "footer": ("helv", 8)}
    doc = pymupdf.open()
    for items in pages:
        page = doc.new_page()
        for style, text, rect in items:
            fontname, size = fonts[style]
            fit = page.insert_textbox(pymupdf.Rect(*rect), text, fontsize=size, fontname=fontname)
            assert fit >= 0, f"fixture text did not fit its box (make the rect taller): {text[:30]!r}"
    doc.save(path)
    doc.close()
    return path


def notes_pdf(path: Path) -> Path:
    def page(n: int, extra: list | None = None):
        items = [("footer", f"Course Notes | Page {n}", (72, 760, 400, 780))]
        return items + (extra or [])

    first = page(1, [
        ("title", "CV Notes", (72, 40, 540, 80)),
        ("h1", "1. Edge Detection", (72, 90, 540, 120)),
        ("body", "Edges are locations of strong intensity change. They outline object boundaries.", (72, 125, 540, 170)),
        ("h2", "Sobel Operator", (72, 185, 540, 215)),
        ("body", "Sobel is a first-order derivative operator that estimates horizontal and vertical gradients.", (72, 220, 540, 270)),
    ])
    second = page(2, [
        ("h1", "2. Thresholding", (72, 90, 540, 120)),
        ("body", "Otsu thresholding picks a threshold automatically from the histogram.", (72, 125, 540, 170)),
    ])
    third = page(3, [
        ("body", "Global thresholding uses a single threshold for the whole image.", (72, 90, 540, 140)),
    ])
    return make_pdf(path, [first, second, third])


def notes_docx(path: Path) -> Path:
    doc = Document()
    doc.add_heading("Unit Notes", 0)
    doc.add_heading("Interpolation", 1)
    doc.add_paragraph("Interpolation estimates pixel values at new locations.")
    doc.add_heading("Bilinear", 2)
    doc.add_paragraph("Uses the four nearest pixels.", style="List Bullet")
    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Method", "Neighbours"
    table.rows[1].cells[0].text, table.rows[1].cells[1].text = "Nearest", "1"
    table.rows[2].cells[0].text, table.rows[2].cells[1].text = "Bicubic", "16"
    doc.save(path)
    return path


def notes_pptx(path: Path) -> Path:
    deck = Presentation()
    layout = deck.slide_layouts[1]  # title + content
    slide = deck.slides.add_slide(layout)
    slide.shapes.title.text = "RANSAC"
    slide.placeholders[1].text_frame.text = "Robust model fitting that rejects outliers."
    slide.notes_slide.notes_text_frame.text = "Mention inliers and outliers."
    slide2 = deck.slides.add_slide(layout)
    slide2.shapes.title.text = "Stitching"
    slide2.placeholders[1].text_frame.text = "Warp overlapping images and blend the seams."
    deck.slides.add_slide(deck.slide_layouts[6])  # blank, no text
    deck.save(path)
    return path
