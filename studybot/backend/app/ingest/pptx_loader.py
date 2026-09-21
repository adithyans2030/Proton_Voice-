"""PowerPoint loader (python-pptx). One segment per slide, including tables and speaker notes."""
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.ingest.text_utils import clean_text, prettify_stem, render_table, sha256_file
from app.ingest.types import IngestError, LoadedDocument, Segment


def _walk(shapes):
    for shape in sorted(shapes, key=lambda s: (s.top or 0, s.left or 0)):
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shape.shapes)
        else:
            yield shape


def load_pptx(path: str | Path) -> LoadedDocument:
    path = Path(path)
    try:
        presentation = Presentation(path)
    except Exception as exc:
        raise IngestError(f"Could not open PowerPoint '{path.name}': {exc}") from exc

    segments: list[Segment] = []
    warnings: list[str] = []

    for number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = clean_text(title_shape.text_frame.text) if title_shape is not None and title_shape.has_text_frame else ""
        parts: list[str] = []
        has_picture = False
        for shape in _walk(slide.shapes):
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                has_picture = True
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if shape.has_text_frame:
                text = clean_text("\n".join("".join(run.text for run in p.runs) for p in shape.text_frame.paragraphs))
                if text:
                    parts.append(text)
            elif getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                table_text = render_table(rows)
                if table_text:
                    parts.append(table_text)
        if slide.has_notes_slide:
            notes = clean_text(slide.notes_slide.notes_text_frame.text)
            if notes:
                parts.append("Speaker notes: " + notes)

        body = "\n".join(parts)
        if not body and not title:
            if has_picture:
                warnings.append(f"Slide {number} contains only images (no text extracted).")
            continue
        label = f"Slide {number}: {title}" if title else f"Slide {number}"
        segments.append(Segment(text=body or title, heading=(label,), slide=number))

    if not segments:
        raise IngestError(f"No text found in '{path.name}'.")

    return LoadedDocument(
        title=prettify_stem(path),
        source_type="pptx",
        source=str(path.resolve()),
        content_hash=sha256_file(path),
        segments=segments,
        warnings=warnings,
    )
